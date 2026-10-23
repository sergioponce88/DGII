"""Generación de documentos Word (python-docx) con el membrete unificado de la D.G.I.I.

Hoja oficio (legal), letra Arial 12, texto justificado con sangría, membrete con escudos
y pie institucional. Incluye utilidades de fechas, horas y números en letras.
"""
import io
import os
from datetime import date, datetime, time, timedelta

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor

AQUI = os.path.dirname(os.path.abspath(__file__))
LOGO_POL = os.path.join(AQUI, "assets", "logo_policia.png")
LOGO_DGII = os.path.join(AQUI, "assets", "logo_dgii.png")
NAVY = RGBColor(0x0F, 0x1E, 0x3C)
GRIS = RGBColor(0x55, 0x5F, 0x6D)
FUENTE = "Arial"
SANGRIA = Cm(3.0)

MESES = ["ENERO", "FEBRERO", "MARZO", "ABRIL", "MAYO", "JUNIO", "JULIO", "AGOSTO", "SEPTIEMBRE", "OCTUBRE", "NOVIEMBRE", "DICIEMBRE"]

# ───────────────────────── Números y fechas en letras ─────────────────────────
_U = ["cero", "uno", "dos", "tres", "cuatro", "cinco", "seis", "siete", "ocho", "nueve", "diez", "once", "doce", "trece",
      "catorce", "quince", "dieciséis", "diecisiete", "dieciocho", "diecinueve", "veinte", "veintiuno", "veintidós",
      "veintitrés", "veinticuatro", "veinticinco", "veintiséis", "veintisiete", "veintiocho", "veintinueve"]
_D = {3: "treinta", 4: "cuarenta", 5: "cincuenta", 6: "sesenta", 7: "setenta", 8: "ochenta", 9: "noventa"}
_C = {1: "ciento", 2: "doscientos", 3: "trescientos", 4: "cuatrocientos", 5: "quinientos", 6: "seiscientos",
      7: "setecientos", 8: "ochocientos", 9: "novecientos"}


def letras(n: int) -> str:
    n = int(n)
    if n < 30:
        return _U[n]
    if n < 100:
        return _D[n // 10] + (f" y {_U[n % 10]}" if n % 10 else "")
    if n < 1000:
        if n == 100:
            return "cien"
        return _C[n // 100] + (f" {letras(n % 100)}" if n % 100 else "")
    if n < 1_000_000:
        m, r = divmod(n, 1000)
        pre = "mil" if m == 1 else f"{_apocope(letras(m))} mil"
        return pre + (f" {letras(r)}" if r else "")
    m, r = divmod(n, 1_000_000)
    pre = "un millón" if m == 1 else f"{_apocope(letras(m))} millones"
    return pre + (f" {letras(r)}" if r else "")


def _apocope(t):
    if t.endswith("veintiuno"):
        return t[:-9] + "veintiún"
    if t.endswith("uno"):
        return t[:-3] + "un"
    return t


def legajo_letras(lp) -> str:
    try:
        return letras(int(str(lp).replace(".", "").strip()))
    except ValueError:
        return str(lp)


def fecha_letras(d: date) -> str:
    """'a los VEINTIDÓS días del mes de OCTUBRE del año DOS MIL VEINTISÉIS'"""
    dia = "al PRIMER día" if d.day == 1 else f"a los {_apocope(letras(d.day)).upper()} días"
    return f"{dia} del mes de {MESES[d.month - 1]} del año {letras(d.year).upper()}"


def hora_letras(h: time) -> str:
    hh = letras(h.hour).upper()
    return f"horas {hh}" + (f" con {letras(h.minute).upper()} minutos" if h.minute else "")


def fecha_larga(d: date) -> str:
    return f"{d.day:02d} de {MESES[d.month - 1].capitalize()} de {d.year}"


def fecha_corta(d) -> str:
    return d.strftime("%d/%m/%Y") if d else "—"


def dias_habiles(desde: date, n: int) -> date:
    d, cuenta = desde, 0
    while True:
        if d.weekday() < 5:
            cuenta += 1
            if cuenta == n:
                return d
        d += timedelta(days=1)


# ───────────────────────── Documento base ─────────────────────────
class Doc:
    def __init__(self, cfg=None, hoja="oficio", membrete=True):
        self.cfg = cfg or {}
        self.d = Document()
        st = self.d.styles["Normal"]
        st.font.name = FUENTE
        st.font.size = Pt(12)
        st.element.rPr.rFonts.set(qn("w:eastAsia"), FUENTE)
        pf = st.paragraph_format
        pf.space_after = Pt(6)
        pf.line_spacing = 1.4
        sec = self.d.sections[0]
        self._hoja(sec, hoja)
        if membrete:
            self._membrete(sec)
            self._pie(sec)

    # ── página ──
    def _hoja(self, sec, hoja):
        sec.page_width = Inches(8.5)
        sec.page_height = Inches(7) if hoja == "media" else Inches(14)
        sec.left_margin, sec.right_margin = Cm(3), Cm(2)
        sec.top_margin = Cm(1.2) if hoja == "media" else Cm(1.5)
        sec.bottom_margin = Cm(1.2) if hoja == "media" else Cm(1.8)
        sec.header_distance, sec.footer_distance = Cm(0.8), Cm(0.8)

    def _membrete(self, sec):
        h = sec.header
        h.is_linked_to_previous = False
        p0 = h.paragraphs[0]
        t = h.add_table(rows=1, cols=3, width=Cm(16.5))
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        fijar_anchos(t, (Cm(2.6), Cm(11.3), Cm(2.6)))
        for c, logo in ((t.rows[0].cells[0], LOGO_POL), (t.rows[0].cells[2], LOGO_DGII)):
            if os.path.exists(logo):
                pr = c.paragraphs[0]
                pr.alignment = WD_ALIGN_PARAGRAPH.CENTER
                pr.add_run().add_picture(logo, height=Cm(2.0))
        mid = t.rows[0].cells[1]
        lineas = [
            ("POLICÍA DE TUCUMÁN", 8.5, True, GRIS),
            ("DIRECCIÓN GENERAL DE INSTITUTOS E INSTRUCCIÓN", 12.5, True, NAVY),
            ("Instituto de Enseñanza Superior de Policía “Grl. José Francisco de San Martín”", 9, False, NAVY),
            (self.cfg.get("doc_direccion") or "Muñecas 1025 · San Miguel de Tucumán · Tel. (0381) 430-5141 · C.P. 4000", 7.5, False, GRIS),
        ]
        mid.paragraphs[0].text = ""
        for i, (txt, sz, b, col) in enumerate(lineas):
            pr = mid.paragraphs[0] if i == 0 else mid.add_paragraph()
            pr.alignment = WD_ALIGN_PARAGRAPH.CENTER
            pr.paragraph_format.space_after = Pt(0)
            pr.paragraph_format.line_spacing = 1.0
            r = pr.add_run(txt)
            r.font.size, r.font.bold, r.font.color.rgb = Pt(sz), b, col
        _celda_vertical(t)
        # línea inferior
        p0.paragraph_format.space_after = Pt(0)
        lin = h.add_paragraph()
        lin.paragraph_format.space_after = Pt(4)
        _borde_inferior(lin, "1F3A6E", 12)

    def _pie(self, sec):
        f = sec.footer
        f.is_linked_to_previous = False
        p = f.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _borde_superior(p, "B9C2D0", 4)
        r = p.add_run(self.cfg.get("doc_lema") or "Dirección General de Institutos e Instrucción · Policía de Tucumán")
        r.font.size, r.font.italic, r.font.color.rgb = Pt(8), True, GRIS

    # ── bloques ──
    def parrafo(self, texto="", sangria=True, alin="j", negrita=False, tam=None, antes=0, despues=6, mayus=False, cursiva=False, color=None):
        p = self.d.add_paragraph()
        p.alignment = {"j": WD_ALIGN_PARAGRAPH.JUSTIFY, "c": WD_ALIGN_PARAGRAPH.CENTER,
                       "r": WD_ALIGN_PARAGRAPH.RIGHT, "l": WD_ALIGN_PARAGRAPH.LEFT}[alin]
        if sangria:
            p.paragraph_format.first_line_indent = SANGRIA
        p.paragraph_format.space_before, p.paragraph_format.space_after = Pt(antes), Pt(despues)
        if texto:
            self.rico(p, texto.upper() if mayus else texto, negrita, tam, cursiva, color)
        return p

    @staticmethod
    def rico(p, texto, negrita=False, tam=None, cursiva=False, color=None):
        """Admite **negrita** dentro del texto."""
        partes = str(texto).split("**")
        for i, t in enumerate(partes):
            if not t:
                continue
            r = p.add_run(t)
            r.font.bold = negrita or (i % 2 == 1)
            r.font.italic = cursiva
            if tam:
                r.font.size = Pt(tam)
            if color is not None:
                r.font.color.rgb = color
        return p

    def fecha(self, d: date, sufijo=".-"):
        self.parrafo(f"San Miguel de Tucumán, {fecha_larga(d)}{sufijo}", sangria=False, alin="r", despues=12)

    def titulo(self, texto, tam=13, subrayar=True, antes=4, despues=10):
        p = self.parrafo(sangria=False, alin="c", antes=antes, despues=despues)
        r = p.add_run(texto)
        r.font.bold, r.font.size, r.font.underline, r.font.color.rgb = True, Pt(tam), subrayar, NAVY
        return p

    def destinatario(self, a: dict):
        """a = {tratamiento, cargo, titular, saludo}"""
        for txt, b in ((a.get("tratamiento") or "Señor", False), (a.get("cargo"), False), (a.get("titular"), True)):
            if txt:
                p = self.parrafo(sangria=False, alin="l", despues=0)
                self.rico(p, txt, negrita=b)
        s = self.parrafo(sangria=False, alin="l", despues=10)
        r = s.add_run(f"Su{' ' * 12}/{' ' * 12}{(a.get('saludo') or 'Despacho').replace('Su ', '')}:")
        r.font.underline = True

    def referencia(self, texto, extra=None):
        """Recuadro sombreado con la referencia del expediente."""
        t = self.d.add_table(rows=1, cols=1)
        fijar_anchos(t, (Cm(16.5),))
        c = t.rows[0].cells[0]
        _sombrear(c, "EEF2F8")
        _borde_izq(c, "1F3A6E", 24)
        p = c.paragraphs[0]
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.line_spacing = 1.15
        self.rico(p, texto, tam=10.5)
        if extra:
            q = c.add_paragraph()
            q.paragraph_format.space_after = Pt(2)
            self.rico(q, extra, tam=9.5, color=GRIS)
        self.d.add_paragraph().paragraph_format.space_after = Pt(2)

    def firma(self, nombre="", cargo="", lado="r", rubrica=True):
        sp = self.d.add_paragraph()
        sp.paragraph_format.space_after, sp.paragraph_format.keep_with_next = Pt(30), True
        al = {"r": WD_ALIGN_PARAGRAPH.RIGHT, "c": WD_ALIGN_PARAGRAPH.CENTER, "l": WD_ALIGN_PARAGRAPH.LEFT}[lado]
        if rubrica:
            p = self.d.add_paragraph("_" * 34)
            p.alignment, p.paragraph_format.space_after = al, Pt(0)
            p.paragraph_format.keep_with_next = True
        for txt, b, sz in ((nombre, True, 11), (cargo, False, 9.5)):
            if txt:
                p = self.d.add_paragraph()
                p.alignment = al
                p.paragraph_format.keep_with_next = True
                p.paragraph_format.space_after, p.paragraph_format.line_spacing = Pt(0), 1.0
                r = p.add_run(txt)
                r.font.bold, r.font.size = b, Pt(sz)

    def firmas_dobles(self, izq, der):
        """Dos firmas lado a lado: [(nombre, cargo)]"""
        self.d.add_paragraph().paragraph_format.space_after = Pt(28)
        sp = self.d.paragraphs[-1]
        sp.paragraph_format.keep_with_next = True
        t = self.d.add_table(rows=1, cols=2)
        fijar_anchos(t, (Cm(8.25), Cm(8.25)))
        trPr = t.rows[0]._tr.get_or_add_trPr()
        cs = OxmlElement("w:cantSplit")
        trPr.append(cs)
        for c, (n, cg) in zip(t.rows[0].cells, (izq, der)):
            p = c.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.add_run("_" * 30 + "\n")
            r = p.add_run(n or "")
            r.font.bold, r.font.size = True, Pt(10)
            if cg:
                r2 = p.add_run("\n" + cg)
                r2.font.size = Pt(9)

    def renglones_firma(self, items):
        """Pie de cédula: líneas para completar a mano."""
        for txt in items:
            p = self.d.add_paragraph()
            p.paragraph_format.space_before, p.paragraph_format.space_after = Pt(14), Pt(0)
            p.add_run("_" * 42)
            q = self.d.add_paragraph()
            q.paragraph_format.space_after = Pt(0)
            r = q.add_run(txt)
            r.font.size, r.font.bold = Pt(8.5), True

    def caja_titulo(self, texto, tam=12):
        t = self.d.add_table(rows=1, cols=1)
        fijar_anchos(t, (Cm(16.5),))
        c = t.rows[0].cells[0]
        _sombrear(c, "1F3A6E")
        p = c.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = p.paragraph_format.space_after = Pt(3)
        r = p.add_run(texto)
        r.font.bold, r.font.size, r.font.color.rgb = True, Pt(tam), RGBColor(0xFF, 0xFF, 0xFF)
        self.d.add_paragraph().paragraph_format.space_after = Pt(0)

    def salto(self):
        self.d.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    def bytes(self) -> bytes:
        buf = io.BytesIO()
        self.d.save(buf)
        return buf.getvalue()


# ───────────────────────── XML auxiliares ─────────────────────────
def fijar_anchos(tabla, anchos):
    """Ancho fijo real (Word y LibreOffice ignoran cell.width si la tabla es autoajustable)."""
    tabla.autofit = False
    tblPr = tabla._tbl.tblPr
    for tag in ("w:tblW", "w:tblLayout"):
        for e in tblPr.findall(qn(tag)):
            tblPr.remove(e)
    w = OxmlElement("w:tblW")
    w.set(qn("w:w"), str(int(sum(a for a in anchos) / 635)))
    w.set(qn("w:type"), "dxa")
    tblPr.append(w)
    lay = OxmlElement("w:tblLayout")
    lay.set(qn("w:type"), "fixed")
    tblPr.append(lay)
    grid = tabla._tbl.tblGrid
    for gc, a in zip(grid.findall(qn("w:gridCol")), anchos):
        gc.set(qn("w:w"), str(int(a / 635)))
    for fila in tabla.rows:
        for c, a in zip(fila.cells, anchos):
            c.width = a


def _sombrear(celda, color):
    tcPr = celda._tc.get_or_add_tcPr()
    sh = OxmlElement("w:shd")
    sh.set(qn("w:val"), "clear")
    sh.set(qn("w:color"), "auto")
    sh.set(qn("w:fill"), color)
    tcPr.append(sh)


def _borde_izq(celda, color, sz):
    tcPr = celda._tc.get_or_add_tcPr()
    b = OxmlElement("w:tcBorders")
    for lado in ("top", "left", "bottom", "right"):
        e = OxmlElement(f"w:{lado}")
        if lado == "left":
            e.set(qn("w:val"), "single")
            e.set(qn("w:sz"), str(sz))
            e.set(qn("w:color"), color)
        else:
            e.set(qn("w:val"), "nil")
        b.append(e)
    tcPr.append(b)


def _celda_vertical(tabla):
    for c in tabla.rows[0].cells:
        tcPr = c._tc.get_or_add_tcPr()
        v = OxmlElement("w:vAlign")
        v.set(qn("w:val"), "center")
        tcPr.append(v)
    tblPr = tabla._tbl.tblPr
    b = OxmlElement("w:tblBorders")
    for lado in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement(f"w:{lado}")
        e.set(qn("w:val"), "nil")
        b.append(e)
    tblPr.append(b)


def _borde_parrafo(p, lado, color, sz):
    pPr = p._p.get_or_add_pPr()
    bdr = OxmlElement("w:pBdr")
    e = OxmlElement(f"w:{lado}")
    e.set(qn("w:val"), "single")
    e.set(qn("w:sz"), str(sz))
    e.set(qn("w:space"), "1")
    e.set(qn("w:color"), color)
    bdr.append(e)
    pPr.append(bdr)


def _borde_inferior(p, color, sz):
    _borde_parrafo(p, "bottom", color, sz)


def _borde_superior(p, color, sz):
    _borde_parrafo(p, "top", color, sz)
